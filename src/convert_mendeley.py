"""Chuyển file xlsx Mendeley (hxh8kd3snj) thành các file wav 16 kHz + metadata.csv.

Layout xlsx (đã kiểm tra, xem data/raw/mendeley/README.md):
  - 3 sheet "Ridge A/B/C", mỗi sheet 127 cột, 1 hàng header + 132 300 mẫu (44,1 kHz, ~3 s).
  - Header: "c<id>_<label>" với label ∈ {im, m, om}; 5 cột c106–c110 ở sheet A/B thiếu tiền tố "c".

Đầu ra:
  data/processed/mendeley/<label>/c<id>_<ridge>.wav   (PCM16 mono, 16 kHz)
  data/processed/mendeley/metadata.csv                 (schema dùng chung với record.py)

Chạy:
  python -m src.convert_mendeley [--xlsx PATH] [--out DIR] [--sr 16000]
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import numpy as np
import soundfile as sf
from scipy.signal import resample_poly

SRC_SR = 44_100
HEADER_RE = re.compile(r"^\s*c?\s*(\d+)\s*[_\-\s]?\s*(im|m|om)\s*$", re.IGNORECASE)
SHEET_TO_RIDGE = {"Ridge A": "A", "Ridge B": "B", "Ridge C": "C"}

# Schema metadata dùng chung cho mọi nguồn dữ liệu (Mendeley + tự thu).
META_FIELDS = ["file", "fruit_id", "ridge", "label", "device", "date", "source", "notes"]


def parse_header(h: object) -> tuple[str, str] | None:
    """'c12_om' hoặc '106_om' -> ('c12', 'om'). Trả về None nếu không khớp."""
    m = HEADER_RE.match(str(h))
    if not m:
        return None
    return f"c{int(m.group(1))}", m.group(2).lower()


def resample(y: np.ndarray, src_sr: int, dst_sr: int) -> np.ndarray:
    """Resample đa pha (polyphase), 44100 -> 16000 tương đương up 160 / down 441."""
    if src_sr == dst_sr:
        return y
    from math import gcd

    g = gcd(src_sr, dst_sr)
    return resample_poly(y, dst_sr // g, src_sr // g).astype(np.float32)


def to_pcm16(y: np.ndarray) -> np.ndarray:
    """Chuẩn hoá peak về 0.9 rồi ép về int16 (tránh clipping, giữ tỉ lệ biên độ trong 1 tín hiệu)."""
    peak = float(np.max(np.abs(y))) if y.size else 0.0
    if peak > 0:
        y = y / peak * 0.9
    return (np.clip(y, -1.0, 1.0) * 32767).astype(np.int16)


def read_sheet(ws, n_samples_hint: int = 132_300) -> tuple[list[tuple[str, str]], np.ndarray]:
    """Đọc một sheet read-only theo hàng, trả về (danh sách (fruit_id, label) theo cột, mảng (n_samples, n_cols))."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    cols: list[tuple[str, str]] = []
    keep_idx: list[int] = []
    for i, h in enumerate(header):
        if h is None:
            continue
        parsed = parse_header(h)
        if parsed is None:
            print(f"  [cảnh báo] bỏ qua cột {i} header không hợp lệ: {h!r}", file=sys.stderr)
            continue
        cols.append(parsed)
        keep_idx.append(i)

    buf = np.zeros((n_samples_hint, len(keep_idx)), dtype=np.float32)
    n = 0
    for r in rows:
        if n >= buf.shape[0]:
            buf = np.vstack([buf, np.zeros((n_samples_hint // 10, buf.shape[1]), dtype=np.float32)])
        buf[n] = [0.0 if (v is None or not isinstance(v, (int, float))) else float(v) for v in (r[i] for i in keep_idx)]
        n += 1
    return cols, buf[:n]


def convert(xlsx: Path, out_dir: Path, dst_sr: int) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows_meta: list[dict] = []
    for sheet in wb.sheetnames:
        ridge = SHEET_TO_RIDGE.get(sheet, sheet.replace(" ", "_"))
        print(f"Đọc sheet {sheet!r} ...", flush=True)
        cols, sig = read_sheet(wb[sheet])
        print(f"  {sig.shape[1]} tín hiệu x {sig.shape[0]} mẫu ({sig.shape[0] / SRC_SR:.2f} s)", flush=True)
        for j, (fruit_id, label) in enumerate(cols):
            y = resample(sig[:, j], SRC_SR, dst_sr)
            rel = Path(label) / f"{fruit_id}_{ridge}.wav"
            dst = out_dir / rel
            dst.parent.mkdir(parents=True, exist_ok=True)
            sf.write(dst, to_pcm16(y), dst_sr, subtype="PCM_16")
            rows_meta.append(
                {
                    "file": rel.as_posix(),
                    "fruit_id": f"mendeley_{fruit_id}",
                    "ridge": ridge,
                    "label": label,
                    "device": "mendeley_atr2500",
                    "date": "2022",
                    "source": "mendeley_hxh8kd3snj_v1",
                    "notes": "",
                }
            )
    return rows_meta


def write_metadata(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=META_FIELDS)
        w.writeheader()
        w.writerows(rows)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", type=Path, default=Path("data/raw/mendeley/coconut_acoustic_signals.xlsx"))
    p.add_argument("--out", type=Path, default=Path("data/processed/mendeley"))
    p.add_argument("--sr", type=int, default=16_000)
    a = p.parse_args(argv)

    if not a.xlsx.exists():
        print(f"Không thấy {a.xlsx}. Xem data/raw/mendeley/README.md để tải.", file=sys.stderr)
        return 1
    a.out.mkdir(parents=True, exist_ok=True)
    rows = convert(a.xlsx, a.out, a.sr)
    write_metadata(rows, a.out / "metadata.csv")

    from collections import Counter

    by_label = Counter(r["label"] for r in rows)
    fruits = Counter(r["label"] for r in {r["fruit_id"]: r for r in rows}.values())
    print(f"\nXong: {len(rows)} wav -> {a.out}")
    print("Tín hiệu / lớp:", dict(by_label))
    print("Quả / lớp     :", dict(fruits))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
